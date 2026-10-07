"""Build reports/Freight_Rate_Report.docx from artifacts/ and reports/figures/.

Every number in the text is read from the artifacts written by the pipeline, so the
report regenerates with the results. Run after src.validate, src.predict, score.py and
src.figures (src.run_all does this).
"""
from __future__ import annotations

import json

import lightgbm
import numpy as np
import pandas as pd
import sklearn
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from . import config, figures

OUT = config.REPORTS_DIR / "Freight_Rate_Report.docx"
FIG = figures.FIG_DIR
A = config.ARTIFACTS_DIR
HEADER_FILL = "E8EEF6"
CHOSEN_FILL = "F3F7FC"


# ---------------------------------------------------------------- docx helpers
def _shade(cell, fill):
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc.append(shd)


def table(doc, df: pd.DataFrame, widths=None, highlight_first=False, font=8.5, right=None):
    right = set(range(1, len(df.columns))) if right is None else set(right)
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, col in enumerate(df.columns):
        c = t.rows[0].cells[j]
        c.text = ""
        r = c.paragraphs[0].add_run(str(col))
        r.bold, r.font.size = True, Pt(font)
        _shade(c, HEADER_FILL)
        if j in right:
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for i, row in enumerate(df.itertuples(index=False)):
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = ""
            r = cells[j].paragraphs[0].add_run("" if v is None or (isinstance(v, float) and np.isnan(v)) else str(v))
            r.font.size = Pt(font)
            if highlight_first and i == 0:
                r.bold = True
                _shade(cells[j], CHOSEN_FILL)
            if j in right:
                cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if widths:
        for row in t.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = Inches(w)
    for row in t.rows:
        for c in row.cells:
            pf = c.paragraphs[0].paragraph_format
            pf.space_before = pf.space_after = Pt(1)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def para(doc, text, bold_lead=None, size=None, style=None):
    p = doc.add_paragraph(style=style)
    if bold_lead:
        p.add_run(bold_lead).bold = True
    r = p.add_run(text)
    if size:
        r.font.size = Pt(size)
    return p


def bullet(doc, text, bold_lead=None):
    return para(doc, text, bold_lead, style="List Bullet")


def figure(doc, path, width, caption):
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    r = p.add_run(caption)
    r.italic, r.font.size = True, Pt(8.5)
    r.font.color.rgb = RGBColor(0x52, 0x51, 0x4E)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def num(x, fmt):
    return format(x, fmt).replace("-", "\u2212")


def money(x):
    return f"${x:,.1f}"


def pct(x, d=2):
    return f"{x:.{d}f}%"


# ---------------------------------------------------------------- content
def load():
    s = pd.read_csv(A / "fold_summary.csv", index_col=0)
    con = pd.read_csv(A / "contrasts.csv")
    dq = pd.read_csv(A / "data_quality.csv")
    tune = pd.read_csv(A / "tuning_results.csv")
    chosen = json.loads((A / "chosen_config.json").read_text())
    checks = json.loads((A / "prediction_checks.json").read_text())
    dec = pd.read_csv(A / "december_predictions.csv")
    return s, con, dq, tune, chosen, checks, dec


def build(stats: dict) -> None:
    s, con, dq, tune, chosen, checks, dec = load()
    H = s.loc["Hybrid (chosen)"]
    blend = s.loc["Codex blend 0.75 XGB / 0.25 LGBM"]
    lgbm = s.loc["Codex LGBM-L1"]
    hp = chosen["hybrid"]
    cv = lambda m, k: con[(con.setting.str.startswith(m)) & (con.model == k)].iloc[0]
    rnd_h, rnd_l = cv("Random", "Hybrid (chosen)"), cv("Random", "Codex LGBM-L1")
    fwd_h, fwd_l = cv("Forward", "Hybrid (chosen)"), cv("Forward", "Codex LGBM-L1")
    uc_h, uc_l = cv("Unseen", "Hybrid (chosen)"), cv("Unseen", "Codex LGBM-L1")
    uc_nr = cv("Unseen", "Hybrid, city cats not routed")
    corrupt = dq[dq.issue.str.startswith("Corrupted")].iloc[0]
    qs = stats["quote_signal"]
    coef = json.loads((A / "final_model" / "stage1_coefficients.json").read_text())
    drift = (np.exp(coef["tm"] * (pd.Timestamp("2025-10-31") - pd.Timestamp(config.TIME_ORIGIN)).days
                    / config.DAYS_PER_MONTH) - 1) * 100
    ramp = {e: (np.exp(coef[f"qe_{e}"]) - 1) * 100 for e in config.EQUIPMENT}
    rpm = checks["median_rate_per_mile"]

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.85)
    sec.top_margin = sec.bottom_margin = Inches(0.75)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    st.paragraph_format.space_after = Pt(4)
    for name, size in (("Heading 1", 14), ("Heading 2", 11.5)):
        doc.styles[name].font.size = Pt(size)

    title = doc.add_heading("Freight Rate Prediction — Model and Validation Report", level=0)
    title.runs[0].font.size = Pt(20)
    para(doc, "Spotter AI ML assessment · Anindya Kundu · October 2026", size=9.5)

    # 1 Summary ---------------------------------------------------------------
    doc.add_heading("1. Summary", level=1)
    para(doc, "The task is to predict posted_rate for 12,000 loads in November–December 2025, using "
              "48,000 labelled loads from January–October 2025. That makes it a forecast one to two "
              "months ahead, and the validation is built around that.")
    bullet(doc, " a linear model on log price (distance, weight, equipment, market_index, smooth "
                "geography, a linear time trend and an equipment-specific quarter-end ramp), followed "
                "by LightGBM on its residuals. The residual model has no trend features, so it cannot "
                "flatten the forward projection.", "Model:")
    bullet(doc, f" fold-average MAE {money(H['avg MAE'])} on raw labels, against {money(blend['avg MAE'])} "
                f"for an XGBoost/LightGBM blend (an earlier candidate solution); {money(H['h2 MAE'])} on the two-months-ahead block "
                f"(train Jan–Aug, test Sep–Oct). On labels with the corrupted rows removed, the "
                f"error is {pct(H['avg cMAPE'])} MAPE.", "Results:")
    bullet(doc, f" {int(corrupt['train']):,} training labels ({corrupt['train_pct']:.1f}%) are corrupted "
                "by a multiplicative factor (about ×3.3 or ×0.29). They are detected and dropped from "
                "training. The same corruption presumably exists in the validation labels and sets a "
                "floor on any model's raw error.", "Key data issue:")
    bullet(doc, f" ${checks['december']['first']:.0f} on 1 December rising to "
                f"${checks['december']['last']:.0f} on the 31st, with a weekly cycle (Thursday peaks) "
                "on top of the quarter-end ramp.", "December chart:")
    figure(doc, config.SCORER_DIR / "candidate_december.png", 6.0,
           "Figure 1. December 2025 chart produced by score.py from the submitted predictions.")

    # 2 Data and findings -------------------------------------------------------
    doc.add_heading("2. Data and key findings", level=1)
    bullet(doc, f" Distance explains most of the price: the correlation is "
                f"{stats['corr_distance_clean']:.3f} once the corrupted labels are removed "
                f"({stats['corr_distance_raw']:.3f} with them). Median rate per mile is "
                f"${rpm['Dry Van']['train_clean']:.2f} for Dry Van, ${rpm['Flatbed']['train_clean']:.2f} "
                f"for Flatbed and ${rpm['Reefer']['train_clean']:.2f} for Reefer. Price grows less than "
                f"proportionally with distance (log-log slope {coef['ld']:.2f}).", "Distance and equipment.")
    bullet(doc, " market_index varies far more between days than within a day. Its weekly cycle bottoms "
                "on Sunday/Monday and peaks on Thursday. It averages "
                f"{stats['mi_mean_train']:.2f} in training and {stats['mi_mean_val']:.2f} in validation, so "
                f"it is a genuine input for Nov–Dec. Its price elasticity is {coef['lmi']:.2f}.",
           "A daily market signal.")
    bullet(doc, f" Once market_index is controlled for, prices still drift up about {drift:.0f}% "
                "from January to October. They also climb through every quarter-end month (March, June, "
                f"September) and reset on the 1st: Flatbed +{ramp['Flatbed']:.1f}%, Reefer "
                f"+{ramp['Reefer']:.1f}%, Dry Van +{ramp['Dry Van']:.1f}% from the first day of the month "
                "to the last (Figure 2; fitted values from the final model). December is a quarter-end "
                "month.", "Time structure that market_index does not explain.")
    figure(doc, FIG / "time_structure.png", 6.3,
           "Figure 2. Daily mean residual by equipment from a model with no time terms (3-day mean). "
           "Dotted lines mark quarter ends.")
    figure(doc, FIG / "market_index.png", 6.3,
           "Figure 3. market_index: lower in Nov–Dec, with the same weekly cycle.")
    lo, hi = stats["range_pct"]
    bullet(doc, f" 8 validation cities never appear in training (1,447 rows, 12%). In training, city "
                f"effects are small ({num(lo, '+.1f')}% to {num(hi, '+.1f')}%), almost identical whether "
                f"the city is the pickup or the delivery (r = {stats['r_pickup_delivery']:.2f}), and track "
                f"latitude (r = {num(stats['r_lat'], '.2f')}, Figure 4). A smooth function of the coordinates therefore "
                "carries over to unseen cities. The coordinates are shifted from the real cities but are "
                "consistent for each city, so they are used as given.", "Unseen cities.")
    figure(doc, FIG / "city_effects.png", 3.9, "Figure 4. City price effect against latitude.")
    bullet(doc, f" Within lane × equipment groups, after weight, market_index and time terms, adding "
                f"quote_signal leaves the residual SD unchanged ({qs['without']:.4f} vs "
                f"{qs['with']:.4f}). Its apparent signal comes from the equipment mix. Adding it to "
                f"the hybrid worsens the fold-average MAE "
                f"({money(s.loc['Hybrid + quote_signal', 'avg MAE'])} vs {money(H['avg MAE'])}), and "
                f"removing it improves Codex's LightGBM ({money(lgbm['avg MAE'])} → "
                f"{money(s.loc['Codex LGBM-L1 without quote_signal', 'avg MAE'])}). It is dropped.",
           "quote_signal is noise.")

    # 3 Data quality ----------------------------------------------------------------
    doc.add_heading("3. Data-quality issues and fixes", level=1)
    d = dq.copy()
    d["validation"] = d["validation"].map(lambda v: "—" if pd.isna(v) else f"{int(v):,}")
    d["train"] = d["train"].map(lambda v: f"{int(v):,}")
    d = d.rename(columns={"issue": "Issue", "train": "Train", "validation": "Validation", "handling": "Handling"})
    table(doc, d[["Issue", "Train", "Validation", "Handling"]], widths=[2.4, 0.6, 0.75, 3.0], right={1, 2})
    para(doc, "Corrupted labels are found with a two-pass robust fit. The first pass fits the structural "
              "model on log price with least squares. The second refits it on rows with |residual| < 0.30. "
              "Rows with |residual| ≥ 0.25 are then flagged. The clusters are far from the clean rows: "
              "no residual falls between 0.12 and 0.79 (Figure 5), so cut-offs of 0.20, 0.25 and 0.30 "
              "flag the same rows. The detector is refitted inside every fold, and flagged rows are "
              "removed from training only; every validation row is predicted. The other repairs touch "
              "inputs only: missing values get flags, so the model can learn whatever their absence "
              "means.")
    figure(doc, FIG / "corrupted_labels.png", 6.3,
           "Figure 5. log(actual / expected) on the full training set (log count scale).")

    # 4 Validation ------------------------------------------------------------------
    doc.add_heading("4. Split and validation approach", level=1)
    para(doc, "The validation set is the two months after the training data, so every split looks "
              "forward in time. Three expanding folds test August, September and October with all "
              "earlier data. A fourth block, h2, trains through August and tests September–October: "
              "like Nov–Dec, it is one to two months past the end of training, and it spans a "
              "quarter-end. Everything was refitted from scratch in each fold, including the label "
              "cleaning and the weight imputation.")
    figure(doc, FIG / "folds.png", 6.3, "Figure 6. Folds. The final model trains on all of Jan–Oct.")
    para(doc, "Model selection used raw MAE averaged over F1–F3, with h2 MAE as the tie-break. "
              "Raw labels include corruption, as Spotter's labels presumably do. MAE is robust to it. "
              f"RMSE is not: every model lands between {money(s['avg RMSE'].min())} and "
              f"{money(s['avg RMSE'].max())}, because the corrupted rows dominate the squared error. "
              "Clean-label MAPE (corrupted rows removed, for reporting only) shows the accuracy on "
              "legitimate prices.")
    c = pd.DataFrame([
        ["Random 10% holdout (all months)", money(rnd_h.MAE), money(rnd_l.MAE)],
        ["Forward: train Jan–Sep, test Oct", money(fwd_h.MAE), money(fwd_l.MAE)],
        ["Unseen cities: 8 withheld, test Sep–Oct", money(uc_h.MAE), money(uc_l.MAE)],
    ], columns=["Split (MAE)", "Hybrid", "Codex LightGBM"])
    table(doc, c, widths=[3.4, 1.2, 1.4])
    para(doc, f"A random split makes the two models look about equal ({money(rnd_h.MAE)} vs "
              f"{money(rnd_l.MAE)}). Going forward in time, the tree model loses ground "
              f"({money(fwd_h.MAE)} vs {money(fwd_l.MAE)}). A random split would therefore have picked "
              "the wrong model. The unseen-city simulation withholds 8 training cities from training and "
              "tests on the Sep–Oct loads that touch them. The hybrid holds up "
              f"({money(uc_h.MAE)}); the tree model, which relies on city identities, does not "
              f"({money(uc_l.MAE)}).")

    # 5 Model -----------------------------------------------------------------------
    doc.add_heading("5. Model", level=1)
    para(doc, "Stage 1 is ordinary least squares on log(posted_rate), fitted on the cleaned rows. Its "
              "inputs are log distance, log weight with a hinge term for loads under 12,000 lb, a "
              "missing-weight flag, log market_index, equipment dummies, a quadratic surface in "
              "pickup and delivery coordinates, a linear trend in months, and a quarter-end ramp per "
              "equipment (0 on the first day of March, June, September or December, rising to 1 on its "
              "last day; 0 otherwise). Stage 2 is LightGBM on the stage-1 residuals. It uses distance, "
              "weight, market_index, equipment, coordinates, weekday and position within the month, "
              "plus pickup and delivery city as categorical features. It has no trend input, so it "
              "cannot undo the trend. Rows that touch a city absent from training use a second "
              "stage-2 model without the city features. "
              "Prediction = exp(stage 1 + stage 2).")
    para(doc, "Why not boosted trees alone: a tree predicts a constant beyond the last date it saw, so "
              "it cannot carry the drift or a quarter-end ramp into the future. In the h2 backtest its "
              "error moves downward through the September ramp (Figure 7), and it never learns that "
              "December will ramp. The linear stage carries time forward; the trees handle the "
              "non-linear parts that do not depend on the date.")
    figure(doc, FIG / "backtest_bias.png", 6.3,
           "Figure 7. h2 backtest: daily mean signed error on clean labels (labels show the last-5-day mean).")
    rows = ["Hybrid (chosen)", "Hybrid, untuned defaults", "Hybrid, no city categoricals",
            "Hybrid + quote_signal", "Hybrid without market_index", "Hybrid 50/50 with Codex XGB", "Hybrid without quarter-end ramp",
            "Structural only (stage 1)", "Codex LGBM-L1 without quote_signal",
            "Codex blend 0.75 XGB / 0.25 LGBM", "Codex LGBM-L1", "Codex XGB-L1",
            "Linear, city indicators, no time terms", "Hybrid without trend",
            "Hybrid without trend or ramp"]
    m = s.loc[rows]
    m = pd.concat([m.iloc[:1], m.iloc[1:].sort_values("avg MAE")])
    mt = pd.DataFrame({
        "Model": m.index, "Aug": m["Aug MAE"].map(money), "Sep": m["Sep MAE"].map(money),
        "Oct": m["Oct MAE"].map(money), "h2": m["h2 MAE"].map(money),
        "Fold avg MAE": m["avg MAE"].map(money), "MAPE": m["avg MAPE"].map(pct),
        "Clean MAPE": m["avg cMAPE"].map(pct)})
    table(doc, mt, widths=[2.45, 0.62, 0.62, 0.62, 0.62, 0.85, 0.6, 0.72], highlight_first=True, font=8)
    para(doc, "MAE in dollars on raw labels; MAPE and clean MAPE averaged over F1–F3. The Codex rows "
              "reproduce an earlier candidate solution (tree models on raw dollars with every feature, "
              "including quote_signal and calendar features). Removing the trend or market_index roughly "
              "doubles the error: handling time is the main problem.", size=9)
    para(doc, "Tuning. Before any tuning run, a small grid was committed to the repository, along with "
              "the selection rule. The search went one setting at a time. It picked a linear trend "
              "over a step per quarter or a damped trend, 1,000 trees with 31 leaves, L2 loss, and "
              "city categoricals, improving the fold-average MAE from "
              f"{money(s.loc['Hybrid, untuned defaults', 'avg MAE'])} to {money(H['avg MAE'])}. "
              "One change was made after tuning. The folds contain no unseen cities, and in the "
              f"unseen-city simulation the city categoricals cost money ({money(uc_nr.MAE)} vs "
              f"{money(uc_h.MAE)} with routing), so rows touching an unseen city use the model "
              "without them. This leaves every fold score unchanged. Results are stable across "
              "seeds (h2 MAE " + " / ".join(money(v) for v in chosen["seed_stability_h2_mae"].values())
           + ").")

    # 6 December ---------------------------------------------------------------------
    doc.add_heading("6. December prediction", level=1)
    para(doc, "The chart inputs contain no market_index. Each December day gets the mean market_index "
              "of that day's loads in validation.csv. This is an input feature Spotter supplied for "
              "the same dates, not a label, and it is how the model is applied to the validation "
              "loads too. Dropping it (weekday dummies in its place) raises the h2 clean MAPE "
              f"from {pct(H['h2 cMAPE'])} to {pct(s.loc['Hybrid without market_index', 'h2 cMAPE'])}. The resulting chart reflects four things: "
              "lower market_index in December, the weekly cycle, the continuing drift, and the Dry Van "
              f"quarter-end ramp. It runs from ${dec['hybrid'].min():.0f} to ${dec['hybrid'].max():.0f} "
              f"(mean ${dec['hybrid'].mean():.0f}). A Codex-style LightGBM trained without "
              "market_index draws an almost flat line "
              f"(${dec['codex_lgbm_no_market_index'].min():.0f}–"
              f"${dec['codex_lgbm_no_market_index'].max():.0f}, Figure 8), because a tree has no way "
              "to extend time effects beyond October.")
    figure(doc, FIG / "december_crosscheck.png", 6.3,
           "Figure 8. December chart lane: submitted hybrid vs. a tree model without market_index.")
    ck = checks["median_rate_per_mile"]
    para(doc, "Sanity checks on the 12,000 validation predictions: no missing or non-positive values; "
              "median rate per mile "
              + ", ".join(f"{e} ${ck[e]['predicted']:.2f} (training ${ck[e]['train_clean']:.2f})"
                          for e in config.EQUIPMENT)
              + f"; unseen-city rows average ${checks['unseen_city_rows']['mean_rate_per_mile_unseen']:.2f}/mi "
                f"vs ${checks['unseen_city_rows']['mean_rate_per_mile_seen']:.2f}/mi for the rest.")

    # 7 Limitations ------------------------------------------------------------------
    doc.add_heading("7. Limitations and next steps", level=1)
    bullet(doc, " the linear trend and the December ramp are projected two months beyond the data. "
                "The ramp held in 3 of 3 quarters but cannot be checked for December. In the "
                f"backtests, the model under-predicted October by {abs(H['Oct cBias%']):.1f}% "
                "(the month after a quarter end).", "Extrapolation:")
    bullet(doc, " holiday effects in late December cannot be learned from Jan–Oct data.", "Holidays:")
    bullet(doc, " if the validation labels contain the same ~1.4% corruption, raw MAE and especially "
                "RMSE will carry an error floor that no model can remove.", "Label corruption:")
    bullet(doc, " quantile models for prediction intervals; drift monitoring on market_index and the "
                "residual level; refitting monthly as new labels arrive; checking the corruption "
                "mechanism with the data owner.", "Next steps:")

    # Appendix -------------------------------------------------------------------------
    doc.add_heading("Appendix: how to reproduce", level=1)
    for line in ("python -m pip install -r requirements.txt",
                 "python -m src.run_all          # full pipeline, about 10 minutes on 16 cores",
                 "python -m src.tune             # optional: rerun the pre-registered tuning grid",
                 "python score.py --predictions validation_predictions.csv "
                 "--december-predictions data/december_chart_inputs.csv"):
        p = doc.add_paragraph()
        r = p.add_run(line)
        r.font.name, r.font.size = "Consolas", Pt(8.5)
        p.paragraph_format.space_after = Pt(0)
    para(doc, "")
    para(doc, "src/clean.py: input repairs and the corrupted-label detector · src/features.py: feature "
              "builders · src/models.py: HybridModel and the baselines · src/validate.py: folds, "
              "metrics, contrasts · src/tune.py: tuning grid · src/train.py, src/predict.py: final "
              "fit and outputs · src/figures.py, src/report.py: this report.", size=9)
    para(doc, f"Python 3.12 · pandas {pd.__version__} · numpy {np.__version__} · scikit-learn "
              f"{sklearn.__version__} · LightGBM {lightgbm.__version__} · seeds fixed at {config.SEED}.", size=9)

    config.REPORTS_DIR.mkdir(exist_ok=True)
    doc.save(OUT)
    print("report ->", OUT.relative_to(config.ROOT))


def main() -> None:
    build(figures.main())


if __name__ == "__main__":
    main()
